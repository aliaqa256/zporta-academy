"""
Python-Docx Document Renderer Adapter.
"""
import logging
from io import BytesIO
from bs4 import BeautifulSoup
from core.media_export.application.ports.outbound.document_renderer_port import DocumentRendererPort
from core.media_export.domain.entities import DocumentExportEntity
from core.media_export.domain.exceptions import RenderingError
from core.media_export.domain.policies import HtmlSanitizationPolicy

logger = logging.getLogger(__name__)


class DocxRendererAdapter(DocumentRendererPort):
    """Renders DOCX Word documents from lesson entities."""

    def render_pdf(self, document: DocumentExportEntity) -> bytes:
        from core.media_export.adapters.outbound.renderers.weasyprint_renderer_adapter import WeasyPrintRendererAdapter
        return WeasyPrintRendererAdapter().render_pdf(document)

    def render_docx(self, document: DocumentExportEntity) -> bytes:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Pt, RGBColor
        except ImportError as e:
            raise RenderingError(f"python-docx is not installed: {e}") from e

        buffer = BytesIO()
        doc = Document()

        # Title
        heading = doc.add_heading(document.title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata
        doc.add_paragraph(f"Created by: {document.author_name}")
        doc.add_paragraph(f"Date: {document.created_at_str}")
        if document.subject_name:
            doc.add_paragraph(f"Subject: {document.subject_name}")
        if document.course_title:
            doc.add_paragraph(f"Course: {document.course_title}")

        doc.add_paragraph()

        # Content parsing
        clean_content = HtmlSanitizationPolicy.clean_html_for_export(document.content_html)
        soup = BeautifulSoup(clean_content, "html.parser")

        for el in soup.find_all(["h1", "h2", "h3", "p", "li", "ul", "ol"]):
            text = el.get_text().strip()
            if not text:
                continue
            if el.name == "h1":
                doc.add_heading(text, level=1)
            elif el.name == "h2":
                doc.add_heading(text, level=2)
            elif el.name == "h3":
                doc.add_heading(text, level=3)
            elif el.name == "li":
                doc.add_paragraph(text, style="List Bullet")
            elif el.name == "p":
                doc.add_paragraph(text)

        if document.video_url:
            p = doc.add_paragraph(f"Video Resource: {document.video_url}")
            if p.runs:
                p.runs[0].font.color.rgb = RGBColor(0, 0, 255)

        doc.save(buffer)
        return buffer.getvalue()
